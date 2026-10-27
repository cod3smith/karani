"""Submission-ready documents (ADR 0020).

The pack was markdown, which no ATS accepts — every application began
with manual conversion, which is a large part of why 121 packs
produced 0 applications. These render the same content to an ATS-plain
.docx and a PDF for humans.
"""
from __future__ import annotations

import pytest

from karani.documents import (
    find_chrome, parse_markdown, render_all, to_docx, to_html,
)

RESUME = """# Kelyn Paul Njeri
**Senior Software Engineer**
Nairobi, Kenya · +254 799 670 419
kelyn.njeri@gmail.com · [github.com/cod3smith](https://github.com/cod3smith)

## Summary

Staff engineer with 8+ years on distributed data platforms.

## Experience

### DataQRL — Co-Founder, Chief Engineer
- Built a 7-module data-governance platform processing 10M+ records daily.
- Cut p99 latency from 1.4s to 180ms by reworking the ingestion path.

## Education

### BSc Computer Science
"""


# --- parsing ---

def test_parses_header_sections_and_bullets():
    doc = parse_markdown(RESUME)
    assert doc.name == "Kelyn Paul Njeri"
    assert any("Nairobi" in c for c in doc.contact)
    titles = [s.title for s in doc.sections]
    assert titles[:3] == ["Summary", "Experience", "Education"]
    experience = doc.sections[1]
    assert sum(1 for b in experience.blocks if b.kind == "bullet") == 2


def test_markup_is_stripped_for_parsers():
    """Asterisks and link syntax end up inside the parsed field value
    if they survive into the .docx."""
    from karani.documents import strip_markup
    assert strip_markup("**Senior** [site](https://x)") == "Senior site"


def test_tolerates_model_output_that_is_not_clean_markdown():
    doc = parse_markdown("just a line\n\nanother one")
    assert doc.sections == [] and doc.name == ""   # no crash


# --- html / ats safety ---

def test_html_is_single_column_and_table_free():
    html = to_html(parse_markdown(RESUME))
    for banned in ("<table", "column-count", "<textbox", "position:absolute"):
        assert banned not in html, f"{banned} breaks ATS parsing"
    assert "Kelyn Paul Njeri" in html and "<h2>EXPERIENCE</h2>".lower() \
        in html.lower() or "Experience" in html


def test_bold_and_links_survive_into_the_pdf_html():
    html = to_html(parse_markdown(RESUME))
    assert "<strong>Senior Software Engineer</strong>" in html
    assert "href='https://github.com/cod3smith'" in html


# --- docx ---

def test_docx_renders_with_standard_headings_and_no_tables():
    pytest.importorskip("docx")
    import io
    from docx import Document as Docx

    data = to_docx(parse_markdown(RESUME))
    assert data and data[:2] == b"PK"          # a real .docx zip
    d = Docx(io.BytesIO(data))
    assert not d.tables, "tables are the classic ATS parse failure"
    styles = {p.style.name for p in d.paragraphs}
    assert "List Bullet" in styles
    assert any(s.startswith("Heading") for s in styles)
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Kelyn Paul Njeri" in text
    assert "10M+ records daily" in text
    assert "**" not in text, "markdown leaked into the document"


# --- pdf ---

@pytest.mark.skipif(not find_chrome(), reason="no Chrome on this machine")
def test_pdf_is_a_real_pdf():
    out = render_all(RESUME, title="Resume")
    assert out["pdf"][:5] == b"%PDF-", "not a PDF"
    assert len(out["pdf"]) > 2000


def test_render_all_degrades_when_chrome_is_missing(monkeypatch):
    """A missing browser must cost the PDF, never the application."""
    import karani.documents.render as r
    monkeypatch.setattr(r, "find_chrome", lambda: None)
    out = r.render_all(RESUME)
    assert "pdf" not in out
    assert "docx" in out                        # the uploadable one survives
