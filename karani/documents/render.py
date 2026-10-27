"""Submission-ready documents: ATS-safe .docx and a PDF (ADR 0020).

Two renderers over one parsed document, following MIT CAPD conventions:
a centered name and contact header, standard section headings, strict
single column, no tables, text boxes, headers/footers or graphics —
those are what resume parsers mangle, turning an experience section
into gibberish without telling anyone.

- **.docx** (python-docx) is the artifact to upload where a parser
  will eat it. Plain paragraphs and list bullets, nothing exotic.
- **.pdf** is rendered from the same model through headless Chrome,
  for the human who opens it. Chrome is detected at run time; with no
  browser the PDF is skipped and the .docx still ships, because a
  missing renderer must not cost you the application.
"""
from __future__ import annotations

import logging
import shutil
import subprocess
import tempfile
from pathlib import Path

from .parse import Document, parse_markdown, strip_markup

log = logging.getLogger(__name__)

_CHROME_CANDIDATES = (
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    "/Applications/Chromium.app/Contents/MacOS/Chromium",
    "google-chrome", "chromium", "chromium-browser",
)

# MIT-format styling: one column, conservative serif-free type, generous
# margins, nothing a parser has to guess at.
_CSS = """
@page { size: Letter; margin: 0.6in 0.7in; }
* { box-sizing: border-box; }
body { font-family: "Helvetica Neue", Helvetica, Arial, sans-serif;
       font-size: 10.5pt; line-height: 1.38; color: #111; margin: 0; }
header { text-align: center; margin-bottom: 14px; }
h1 { font-size: 17pt; letter-spacing: 0.4px; margin: 0 0 4px; }
.contact { font-size: 9.5pt; color: #333; }
.contact div { margin: 1px 0; }
h2 { font-size: 11pt; text-transform: uppercase; letter-spacing: 0.8px;
     border-bottom: 1px solid #444; padding-bottom: 2px;
     margin: 15px 0 7px; }
h3 { font-size: 10.5pt; margin: 9px 0 3px; }
p { margin: 0 0 6px; }
ul { margin: 0 0 7px; padding-left: 17px; }
li { margin: 0 0 3px; }
a { color: #111; text-decoration: none; }
"""


def find_chrome() -> str | None:
    for candidate in _CHROME_CANDIDATES:
        if candidate.startswith("/") and Path(candidate).exists():
            return candidate
        found = shutil.which(candidate)
        if found:
            return found
    return None


def to_html(doc: Document, title: str = "") -> str:
    out = ["<!doctype html><html><head><meta charset='utf-8'>",
           f"<title>{title or doc.name}</title>",
           f"<style>{_CSS}</style></head><body>", "<header>"]
    if doc.name:
        out.append(f"<h1>{doc.name}</h1>")
    if doc.contact:
        out.append("<div class='contact'>")
        out += [f"<div>{line}</div>" for line in doc.contact]
        out.append("</div>")
    out.append("</header>")

    for section in doc.sections:
        out.append(f"<h2>{section.title}</h2>")
        bullets: list[str] = []
        for block in section.blocks:
            if block.kind == "bullet":
                bullets.append(block.text)
                continue
            if bullets:
                out.append("<ul>" + "".join(f"<li>{b}</li>" for b in bullets)
                           + "</ul>")
                bullets = []
            if block.kind == "subhead":
                out.append(f"<h3>{block.text}</h3>")
                continue
            out.append(f"<p>{block.text}</p>")
        if bullets:
            out.append("<ul>" + "".join(f"<li>{b}</li>" for b in bullets)
                       + "</ul>")
    out.append("</body></html>")
    return _inline_markup("\n".join(out))


def _inline_markup(html: str) -> str:
    """**bold** and [text](url) survive into the PDF as real markup."""
    import re
    html = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", html)
    return re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"<a href='\2'>\1</a>", html)


def to_pdf(doc: Document, title: str = "") -> bytes | None:
    """Render through headless Chrome. None when no browser is present."""
    chrome = find_chrome()
    if not chrome:
        log.warning("no Chrome/Chromium found — skipping PDF, .docx still "
                    "rendered")
        return None
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "doc.html"
        dst = Path(tmp) / "doc.pdf"
        src.write_text(to_html(doc, title), encoding="utf-8")
        try:
            subprocess.run(
                [chrome, "--headless", "--disable-gpu", "--no-sandbox",
                 "--no-pdf-header-footer", f"--print-to-pdf={dst}",
                 src.as_uri()],
                check=True, capture_output=True, timeout=90,
            )
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as e:
            log.warning("chrome PDF render failed: %s", e)
            return None
        return dst.read_bytes() if dst.exists() else None


def to_docx(doc: Document) -> bytes | None:
    """ATS-plain .docx: headings and bullets, no tables or text boxes."""
    try:
        from docx import Document as Docx
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
    except ImportError:
        log.warning("python-docx not installed (`uv sync --extra documents`)"
                    " — skipping .docx")
        return None
    import io

    d = Docx()
    for section in d.sections:          # MIT margins
        section.top_margin = section.bottom_margin = Pt(43)
        section.left_margin = section.right_margin = Pt(50)

    normal = d.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    if doc.name:
        head = d.add_paragraph()
        head.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = head.add_run(doc.name)
        run.bold = True
        run.font.size = Pt(16)
    for line in doc.contact:
        para = d.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.add_run(strip_markup(line)).font.size = Pt(9.5)

    for section in doc.sections:
        # Built-in heading styles: what parsers look for.
        d.add_heading(section.title.upper(), level=1)
        for block in section.blocks:
            text = strip_markup(block.text)
            if not text:
                continue
            if block.kind == "bullet":
                d.add_paragraph(text, style="List Bullet")
            elif block.kind == "subhead":
                d.add_heading(text, level=2)
            else:
                d.add_paragraph(text)

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def render_all(markdown: str, *, title: str = "") -> dict[str, bytes]:
    """{extension: bytes} for whatever this machine can produce."""
    doc = parse_markdown(markdown)
    out: dict[str, bytes] = {}
    pdf = to_pdf(doc, title)
    if pdf:
        out["pdf"] = pdf
    docx = to_docx(doc)
    if docx:
        out["docx"] = docx
    return out
